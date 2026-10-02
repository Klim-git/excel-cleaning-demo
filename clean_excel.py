"""Merge XLSX contact lists conservatively. No network access; sources stay intact."""
import argparse
import csv
import json
import re
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FIELDS = ['Организация', 'Город', 'Телефон', 'Email']

def text(value):
    return re.sub(r'\s+', ' ', str(value or '').strip())

def normalize(row):
    result = {key: text(row.get(key)) for key in FIELDS}
    issues = []
    phone = result['Телефон']
    digits = re.sub(r'\D', '', phone)
    if phone:
        if len(digits) == 11 and digits[0] in '78':
            result['Телефон'] = '+7' + digits[1:]
        elif len(digits) == 10:
            result['Телефон'] = '+7' + digits
        else:
            issues.append('Некорректный телефон')
    email = result['Email'].lower()
    result['Email'] = email
    if email and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+', email):
        issues.append('Некорректный email')
    if not result['Организация']:
        issues.append('Нет организации')
    if not result['Город']:
        issues.append('Нет города')
    if not phone and not email:
        issues.append('Нет контактов')
    return result, issues

def safe(value):
    # Prevent source strings from becoming Excel/CSV formulas.
    return "'" + value if isinstance(value, str) and value.startswith('=') else value

def style(sheet):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = 'C2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        cell.fill = PatternFill('solid', fgColor='183153')
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 45 if sheet.title == 'Дубли' else 32
        for cell in row:
            cell.font = Font(name='Arial', size=11, color='183153')
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    for i in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(i)].width = 28 if i < 6 else 45

def run(input_dir, output_dir):
    input_dir, output_dir = Path(input_dir).resolve(), Path(output_dir).resolve()
    if input_dir == output_dir or input_dir in output_dir.parents:
        raise ValueError('Output must be outside the input directory')
    files = sorted(input_dir.glob('*.xlsx'))
    if not files:
        raise ValueError('No XLSX source files found')
    raw, clean, duplicates, seen = [], [], [], {}
    changes = 0
    for path in files:
        if path.name.startswith('~$'):
            continue
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = wb.active
            iterator = sheet.iter_rows(values_only=True)
            headers = [text(x) for x in next(iterator, ())]
            if any(headers.count(k) != 1 for k in FIELDS):
                raise ValueError(f'{path.name}: required unique columns: {FIELDS}')
            for number, values in enumerate(iterator, 2):
                if not any(v is not None and text(v) for v in values):
                    continue
                row = dict(zip(headers, values))
                original = {k: row.get(k) for k in FIELDS}
                raw.append([path.name, number] + [safe(original[k]) for k in FIELDS])
                normalized, issues = normalize(original)
                if any(str(original[k] if original[k] is not None else '') != normalized[k] for k in FIELDS):
                    changes += 1
                # Only exact normalized full rows with identity and a contact are removed.
                # Same phone/email alone may identify different branches: preserve them.
                key = tuple(normalized[k].casefold() for k in FIELDS)
                eligible = bool(normalized['Организация'] and normalized['Город'] and
                                (normalized['Телефон'] or normalized['Email'])) and not issues
                source = f'{path.name}:{number}'
                if eligible and key in seen:
                    duplicates.append([source, seen[key], 'Полное совпадение после нормализации'] + [safe(normalized[k]) for k in FIELDS])
                    continue
                if eligible:
                    seen[key] = source
                clean.append([safe(normalized[k]) for k in FIELDS] + [source, '; '.join(issues)])
        finally:
            wb.close()
    output_dir.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    report = book.active
    report.title = 'Отчёт'
    metrics = {'Исходных строк': len(raw), 'Сохранено строк': len(clean),
               'Удалено точных дублей': len(duplicates), 'Нормализовано исходных строк': changes,
               'Строк требуют проверки': sum(bool(r[-1]) for r in clean)}
    report.append(['Объединение и очистка Excel', 'Демонстрационные данные'])
    report.append(['Показатель', 'Значение'])
    for name, value in metrics.items():
        report.append([name, value])
    report.append(['Баланс строк', '=B3-B4-B5'])
    report.append(['Правило дублей', 'Только полное нормализованное совпадение всех четырёх полей'])
    report.append(['Проверка email', 'Только формат, без проверки доставки или существования ящика'])
    report.append(['Пропуски и ошибки', 'Сохранены для ручной проверки; контакты не придуманы'])
    report.append(['Сводка', 'Это снимок обработки. Для новых исходников запустите скрипт повторно.'])
    for name, headers, rows in [
        ('Результат', FIELDS + ['Источник', 'Замечания'], clean),
        ('Дубли', ['Удалённая строка', 'Сохранённая строка', 'Причина'] + FIELDS, duplicates),
        ('Исходные строки', ['Файл', 'Строка'] + FIELDS, raw),
    ]:
        sheet = book.create_sheet(name)
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        style(sheet)
    style(report)
    report.freeze_panes = None
    report.auto_filter.ref = None
    report.column_dimensions['A'].width = 38
    report.column_dimensions['B'].width = 85
    book.save(output_dir / 'cleaned.xlsx')
    with (output_dir / 'cleaned.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream, delimiter=';')
        writer.writerow(FIELDS + ['Источник', 'Замечания'])
        writer.writerows([
            ["'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value for value in row]
            for row in clean
        ])
    (output_dir / 'report.json').write_text(json.dumps(metrics, ensure_ascii=False, indent=2), encoding='utf-8')
    assert len(raw) == len(clean) + len(duplicates)
    return metrics

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, default=Path(__file__).parent / 'input')
    parser.add_argument('--output', type=Path, default=Path(__file__).parent / 'output')
    args = parser.parse_args()
    print(json.dumps(run(args.input, args.output), ensure_ascii=False, indent=2))
