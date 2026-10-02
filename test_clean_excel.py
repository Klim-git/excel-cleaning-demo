import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from clean_excel import FIELDS, run

class CleaningTests(unittest.TestCase):
    def test_preserves_branches_errors_and_formula_strings(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / 'input'
            source.mkdir()
            book = Workbook()
            sheet = book.active
            sheet.append(FIELDS)
            sheet.append(['A', 'Уфа', '8 (000) 000-00-01', 'A@EXAMPLE.TEST'])
            sheet.append([' A ', 'Уфа', '+70000000001', 'a@example.test'])
            sheet.append(['A филиал', 'Уфа', '+70000000001', 'a@example.test'])
            sheet.append(['Нет контактов', 'Уфа', '', ''])
            sheet.append(['=1+1', 'Уфа', '123', 'bad'])
            sheet.append([None, None, None, None])
            path = source / 'source.xlsx'
            book.save(path)
            before = path.read_bytes()
            metrics = run(source, root / 'output')
            self.assertEqual(metrics['Исходных строк'], 5)
            self.assertEqual(metrics['Сохранено строк'], 4)
            self.assertEqual(metrics['Удалено точных дублей'], 1)
            self.assertEqual(metrics['Строк требуют проверки'], 2)
            self.assertEqual(path.read_bytes(), before)
            result = load_workbook(root / 'output' / 'cleaned.xlsx')
            self.assertNotEqual(result['Результат']['A5'].data_type, 'f')
            self.assertEqual(result['Дубли']['B2'].value, 'source.xlsx:2')
            result.close()

    def test_bad_headers_and_empty_input_fail(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            with self.assertRaises(ValueError):
                run(root, root.parent / 'unused-result')
            book = Workbook()
            book.active.append(['Организация'])
            book.save(root / 'bad.xlsx')
            with self.assertRaises(ValueError):
                run(root, root.parent / 'unused-result')

    def test_demo_expected_totals_and_repeatability(self):
        root = Path(__file__).parent
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary)
            first = run(root / 'input', destination)
            second = run(root / 'input', destination)
            self.assertEqual(first, second)
            self.assertEqual(first['Исходных строк'], 18)
            self.assertEqual(first['Сохранено строк'], 15)
            self.assertEqual(first['Удалено точных дублей'], 3)
            self.assertEqual(first['Строк требуют проверки'], 4)

if __name__ == '__main__':
    unittest.main()
